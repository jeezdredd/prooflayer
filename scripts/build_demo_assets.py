import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from fpdf import FPDF

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "backend/classroom/demo"
FONT = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")


def build_docx(path: Path, text: str, hidden: str) -> None:
    doc = Document()
    for block in text.split("\n\n"):
        doc.add_paragraph(block.strip())
    if hidden:
        run = doc.add_paragraph().add_run(hidden)
        run.font.size = Pt(1)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.save(path)


def build_pdf(path: Path, text: str) -> None:
    pdf = FPDF(format="A4")
    pdf.set_margins(22, 22, 22)
    pdf.add_page()
    pdf.add_font("ArialUnicode", fname=str(FONT))
    pdf.set_font("ArialUnicode", size=12)
    for block in text.split("\n\n"):
        pdf.multi_cell(0, 7, block.strip())
        pdf.ln(4)
    pdf.output(str(path))


def main() -> None:
    data = json.loads((DEMO / "demo_data.json").read_text(encoding="utf-8"))
    built = []
    for org in data["organizations"]:
        for classroom in org["classes"]:
            for student in classroom["students"]:
                for work in student["works"]:
                    asset = work.get("asset", "")
                    target = DEMO / "assets" / asset
                    if asset.endswith(".docx"):
                        build_docx(target, work["text"], work.get("hidden_text", ""))
                    elif asset.endswith(".pdf"):
                        build_pdf(target, work["text"])
                    else:
                        continue
                    built.append(f"{asset} {target.stat().st_size} bytes")
    print("\n".join(built))


if __name__ == "__main__":
    sys.exit(main())
