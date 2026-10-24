import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from analyzers.aggregator import aggregate
from analyzers.implementations.text_signals import (
    CONTEXT_FLAG,
    TextEvasionAnalyzer,
    TextGradeFitAnalyzer,
    TextStyleAnalyzer,
)
from analyzers.models import AnalysisResult
from analyzers.text.evasion import scan_docx, scan_text
from analyzers.text.grade import assess, lexeme_zipf, readability_band
from analyzers.text.language import detect_language
from analyzers.text.style import drift, features, phrasebook
from classroom.models import Classroom, ShadowSet
from content.tests.factories import SubmissionFactory
from users.tests.factories import UserFactory

from .factories import AnalysisResultFactory, AnalyzerConfigFactory

KID_RU = (
    "Летом я ездил к бабушке в деревню. Там было очень весело. Мы с братом ходили на речку и ловили рыбу. "
    "Один раз мы поймали большую щуку. Бабушка пекла пироги с яблоками, а дедушка показывал нам, как чинить "
    "велосипед. Вечером мы сидели у костра и смотрели на звёзды. Я очень люблю деревню, потому что там можно "
    "бегать весь день и никто не ругается. Ещё у бабушки есть собака Шарик, он очень добрый и любит играть с "
    "мячом. Мне было грустно уезжать домой, но я обещал приехать на следующее лето."
)
ACADEMIC_RU = (
    "Декомпозиция сложной задачи представляет собой ключевой методологический инструмент, позволяющий "
    "структурировать когнитивную нагрузку и оптимизировать процесс принятия решений. Данный подход предполагает "
    "последовательное выделение функциональных компонентов, каждый из которых характеризуется относительной "
    "автономностью и интерпретируемостью. Таким образом, декомпозиция обеспечивает синергетический эффект, "
    "повышая эффективность аналитической деятельности. Кроме того, применение данного инструмента способствует "
    "формированию системного мышления, что является фундаментальной предпосылкой успешной адаптации в условиях "
    "неопределённости. В заключение следует подчеркнуть, что интеграция декомпозиции в образовательный процесс "
    "представляется целесообразной и перспективной стратегией развития метакогнитивных компетенций обучающихся."
)
ADULT_NARRATIVE_RU = (
    "Прошлым летом, когда я гостил у бабушки в деревне, мне довелось впервые в жизни отправиться на утреннюю рыбалку, "
    "и это событие, признаться, произвело на меня неизгладимое впечатление. Мы вышли задолго до рассвета, когда над рекой ещё "
    "стелился густой, почти осязаемый туман, а прибрежные ивы казались причудливыми силуэтами, словно сошедшими со страниц старинной "
    "сказки. Бабушка, обладающая удивительным терпением и поистине энциклопедическими познаниями о повадках речной рыбы, "
    "методично объясняла мне тонкости выбора наживки, особенности течения и закономерности поведения щуки в утренние часы. "
    "Спустя несколько часов напряжённого ожидания мне всё-таки удалось вытащить небольшую, но весьма упитанную щуку, и в тот "
    "момент я испытал непередаваемое чувство гордости, смешанное с благодарностью за этот неспешный, наполненный смыслом день."
)
KID_EN = (
    "Last summer I went to my grandma house in the village. It was really fun. Me and my brother went to the "
    "river every day and we caught fish. One time we caught a big one and grandma cooked it for dinner. My "
    "grandpa showed us how to fix a bike and we rode it all around the field. At night we sat by the fire and "
    "looked at the stars. I love the village because you can run all day and nobody gets mad at you. Grandma "
    "also has a dog named Rex and he likes to play with a ball in the garden. I was sad to go home but I "
    "promised to come back next summer and help grandpa paint the old barn."
)
GRAMMAR_RU = (
    "На уроке русского языка мы повторяли обособленные определения и однородные члены предложения. "
    "Учительница объяснила, что деепричастный оборот всегда выделяется запятыми с двух сторон. "
    "Потом мы искали в стихотворении гиперболы, эпитеты и олицетворения. "
    "Олицетворение помогает автору оживить природу, а гипербола усиливает впечатление от картины. "
    "Я записал правило в тетрадь и придумал свои примеры с однородными членами. "
    "Дома я ещё раз прочитал параграф и сделал упражнение с обособленным определением. "
    "Мне понравилось, что примеры были из знакомых книг, поэтому запоминать правило было легко."
)
OBLIQUE_RU = (
    "Бабушка часто рассказывает про индустриализацию и про то, как её родители строили завод. "
    "Она говорит, что люди тогда отличались дружелюбием и нравственностью, хотя жили очень трудно. "
    "Я иногда устаю, но не жалуюсь на перегруженностью уроками, потому что бабушке было тяжелее. "
    "Мама учит меня запоминанию стихов: мы читаем строчку и повторяем её с закрытыми глазами. "
    "Вчера мы разбирали стихотворение с гиперболой и предложения с однородными членами. "
    "Папа сказал, что хорошая память пригодится мне на экзаменах и в жизни."
)
BOOKISH_RU = (
    "Онегин кажется мне человеком, который сам не знает, чего хочет. Он умный и образованный, но ему всё скучно. "
    "В книге его называют лишним человеком, и я понимаю почему: он не находит себе дела. "
    "Учительница сказала, что это рефлексирующий герой байронического типа, который всё время думает о себе. "
    "Мне его немного жалко, потому что он сам разрушил свою жизнь, когда отказал Татьяне. "
    "Потом он убил друга на дуэли и уехал путешествовать, но и это не помогло. "
    "В конце он понимает, что любит Татьяну, но уже поздно, и она остаётся верной мужу."
)
KID_FR = (
    "L'été dernier je suis allé chez ma grand-mère à la campagne. C'était super. Avec mon frère on allait à la "
    "rivière tous les jours et on pêchait des poissons. Une fois on a attrapé un gros poisson et mamie l'a "
    "cuisiné pour le dîner. Mon grand-père nous a montré comment réparer un vélo et on a roulé partout dans le "
    "champ. Le soir on regardait les étoiles près du feu. J'adore la campagne parce qu'on peut courir toute la "
    "journée et personne ne crie. Mamie a aussi un chien qui s'appelle Max et il aime jouer avec la balle dans "
    "le jardin. J'étais triste de rentrer à la maison mais je vais revenir l'été prochain."
)


class TestLanguage:
    @pytest.mark.parametrize("text,code", [(KID_RU, "ru"), (KID_EN, "en"), (KID_FR, "fr")])
    def test_detects_supported_languages(self, text, code):
        assert detect_language(text).code == code

    def test_too_short_is_unknown(self):
        assert detect_language("Привет, мир") is None


class TestEvasion:
    def test_latin_homoglyphs_inside_russian_words(self):
        report = scan_text("Этот текст нaписан с лaтинской a внутри слов, чтобы обмaнуть систему прoверки.")
        assert len(report.mixed_script_words) == 4
        assert report.strength() >= 0.9

    def test_cyrillic_homoglyphs_inside_english_words(self):
        assert scan_text("This tеxt usеs Cyrillic е in English wоrds.").strength() >= 0.9

    def test_zero_width_spaces(self):
        report = scan_text("Plain​ text​ with​ zero​ width​ spaces.")
        assert report.invisible_total == 5
        assert report.strength() >= 0.9

    def test_evidence_shows_where_hidden_characters_sit(self):
        evidence = scan_text("Исто\u200bрия важна. Один\u2003два и\u2003три. Конец\u2024 Ещё\u2024 Всё\u2024").as_evidence()
        assert evidence["invisible_words"] == ["Исто\u00b7рия"]
        assert evidence["odd_space_words"] == ["Один\u00b7два", "и\u00b7три"]
        assert evidence["punct_words"] == ["Конец\u2024", "Ещё\u2024", "Всё\u2024"]

    def test_evidence_lists_words_with_stray_marks(self):
        evidence = scan_text("The ma\u0323rket i\u0323s gro\u0323wing fast.").as_evidence()
        assert evidence["stray_mark_words"] == ["ma\u0323rket", "i\u0323s", "gro\u0323wing"]

    def test_examples_are_capped(self):
        evidence = scan_text(" ".join(f"сло\u200bво{i}" for i in range(30))).as_evidence()
        assert len(evidence["invisible_words"]) == 10

    def test_lookalike_punctuation(self):
        report = scan_text("The essay ends here\u2024 Another sentence\uff0c with a comma\u2024 And one more\u2024")
        assert sum(report.lookalike_punctuation.values()) == 4
        assert report.strength() >= 0.9

    def test_stray_diacritics_from_a_humanizer(self):
        report = scan_text("The ma\u0323rket i\u0323s gro\u0323wing fast.")
        assert report.stray_marks == 3
        assert report.strength() >= 0.9

    def test_styled_letters_inside_words(self):
        report = scan_text("A w\U0001d428rd and an\uff4fther one.")
        assert len(report.styled_letter_words) == 2
        assert report.strength() >= 0.9

    @pytest.mark.parametrize("text", [
        "Ру\u0301сский язык и ударе\u0301ния в учебнике, ёлка и йод.",
        "Cafe\u0301 cre\u0300me bru\u0302le\u0301e, nai\u0308ve garc\u0327on, c\u030ce\u0301ski.",
        "The \ufb01rst e\ufb03cient \ufb02ow from a PDF.",
        "中文文本，这是测试。很好！" * 5,
    ])
    def test_ordinary_diacritics_ligatures_and_cjk_are_clean(self, text):
        assert scan_text(text).strength() == 0.0

    @pytest.mark.parametrize("text", [KID_RU, KID_EN, KID_FR, "IT-компания и Wi-Fi в школе 2024 года.", "﻿Текст с BOM в начале файла."])
    def test_clean_text_has_no_findings(self, text):
        assert scan_text(text).strength() == 0.0

    def test_docx_hidden_white_and_tiny_runs(self, tmp_path):
        doc = Document()
        paragraph = doc.add_paragraph("Обычный текст работы. ")
        hidden = paragraph.add_run("скрытые слова для обмана антиплагиата")
        hidden.font.hidden = True
        white = paragraph.add_run("белый текст между словами")
        white.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tiny = paragraph.add_run("микрошрифт")
        tiny.font.size = Pt(1)
        path = tmp_path / "essay.docx"
        doc.save(path)
        reasons = {tuple(r["reasons"]) for r in scan_docx(str(path))}
        assert reasons == {("hidden",), ("white text",), ("tiny font",)}


class TestGradeFit:
    def test_academic_text_far_above_eighth_grade(self):
        fit = assess(ACADEMIC_RU, detect_language(ACADEMIC_RU), 8)
        assert fit.mismatch >= 0.6
        assert "декомпозиция" in fit.rare_words

    def test_ordinary_words_in_rare_forms_are_not_beyond_level(self):
        text = (
            "Запоминание новых слов даётся мне легко, если повторять их каждый день понемногу. "
            "Гармоничное расписание помогает избежать перегруженности перед контрольными работами. "
            "Мама говорит, что важно отвечать аргументированно и спокойно, без спешки и лишних слов. "
            "Ещё я заметила, что лучше учить вечером, а утром быстро проверять себя по карточкам. "
            "Так получается запомнить больше и не уставать, даже когда уроков очень много."
        )
        fit = assess(text, detect_language(text), 8)
        assert not {"запоминание", "гармоничное", "перегруженности", "аргументированно"} & set(fit.rare_words)

    @pytest.mark.parametrize("level", [8, 9, 11])
    def test_ordinary_words_in_oblique_cases_are_not_beyond_level(self, level):
        fit = assess(OBLIQUE_RU, detect_language(OBLIQUE_RU), level)
        assert fit.rare_words == []
        assert fit.mismatch < CONTEXT_FLAG

    @pytest.mark.parametrize("level", [8, 9, 11])
    def test_plain_grammar_lesson_stays_below_the_flag(self, level):
        fit = assess(GRAMMAR_RU, detect_language(GRAMMAR_RU), level)
        assert fit.rare_words == []
        assert fit.mismatch < CONTEXT_FLAG

    def test_word_frequency_counts_the_whole_word_not_one_form(self):
        assert lexeme_zipf("запоминанию") == lexeme_zipf("запоминание")
        assert lexeme_zipf("однородными") > 3.0
        assert lexeme_zipf("детерминированность") < 2.0

    def test_look_alike_letters_do_not_count_as_rare_words(self):
        fit = assess(ACADEMIC_RU.replace("о", "o"), detect_language(ACADEMIC_RU), 8)
        assert all("o" not in w for w in fit.rare_words)
        assert fit.as_evidence()["readability_band"] == "adult"

    def test_readability_is_reported_as_plain_facts(self):
        evidence = assess(ACADEMIC_RU, detect_language(ACADEMIC_RU), 8).as_evidence()
        assert evidence["words_per_sentence"] > 15
        assert evidence["syllables_per_word"] > 3
        kid = assess(KID_RU, detect_language(KID_RU), 8).as_evidence()
        assert kid["words_per_sentence"] < 12 and kid["syllables_per_word"] < 2.5
        assert kid["readability_band"] == "plain"

    def test_russian_bands_never_call_a_school_essay_university(self):
        assert readability_band(12.0, "ru") == "textbook"
        assert readability_band(16.8, "ru") == "adult"
        assert readability_band(12.0) == "university"

    def test_same_text_is_fine_for_a_university_student(self):
        assert assess(ACADEMIC_RU, detect_language(ACADEMIC_RU), 13).mismatch < 0.6

    @pytest.mark.parametrize("text", [KID_RU, KID_EN, KID_FR])
    def test_child_writing_fits_fifth_grade(self, text):
        fit = assess(text, detect_language(text), 5)
        assert fit is not None
        assert fit.mismatch < 0.3


class TestStyleDrift:
    def test_sudden_academic_style_drifts_from_child_history(self):
        lang = detect_language(KID_RU)
        history = [features(KID_RU, lang), features(KID_RU.replace("бабушке", "дедушке"), lang), features(KID_RU.replace("щуку", "окуня"), lang)]
        assert drift(features(ACADEMIC_RU, lang), history)["score"] >= 0.6

    def test_similar_work_does_not_drift(self):
        lang = detect_language(KID_RU)
        history = [features(KID_RU, lang), features(KID_RU.replace("бабушке", "дедушке"), lang)]
        assert drift(features(KID_RU.replace("щуку", "окуня"), lang), history)["score"] < 0.3

    def test_drift_says_what_changed(self):
        lang = detect_language(KID_RU)
        history = [features(KID_RU, lang), features(KID_RU.replace("бабушке", "дедушке"), lang), features(KID_RU.replace("щуку", "окуня"), lang)]
        changes = drift(features(ACADEMIC_RU, lang), history)["changes"]
        keys = [c["key"] for c in changes]
        assert "word_length" in keys and len(changes) <= 3
        word_length = next(c for c in changes if c["key"] == "word_length")
        assert word_length["now"] > word_length["usual"]


class TestPhrasebook:
    def test_school_connectors_are_not_chatbot_phrases(self):
        text = "Таким образом, кроме того, в заключение, подводя итог, в целом, более того, в свою очередь, книги важны."
        assert phrasebook(text, detect_language(KID_RU))["hits"] == {}

    def test_chips_show_the_word_from_the_text(self):
        text = "Немаловажную роль играет чтение, ведь оно многогранное. Чтение представляет собой труд."
        hits = phrasebook(text, detect_language(KID_RU))["hits"]
        assert hits == {"немаловажную": 1, "многогранное": 1, "представляет собой": 1}

    def test_words_from_the_assignment_are_dropped(self):
        text = "В современном мире книги важны. Чтение представляет собой труд, и в современном мире это ценно."
        lang = detect_language(KID_RU)
        hits = phrasebook(text, lang, ("Как читать в современном мире?",))["hits"]
        assert hits == {"представляет собой": 1}


def _text_file(tmp_path, text, name="work.txt"):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return str(path)


@pytest.mark.django_db
class TestTextAnalyzers:
    def _class_submission(self, level=8, user=None):
        teacher = UserFactory()
        classroom = Classroom.create_for(teacher, "History 8A", level=level)
        return SubmissionFactory(user=user or UserFactory(), classroom=classroom, mime_type="text/plain")

    def _grade_history(self, student, mismatches):
        config = AnalyzerConfigFactory(name="text_grade_fit")
        for m in mismatches:
            AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence={"language": "ru", "mismatch": m})

    def test_grade_fit_waits_for_enough_earlier_work(self, tmp_path):
        student = UserFactory()
        self._grade_history(student, [0.1, 0.1])
        sub = self._class_submission(level=8, user=student)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, ACADEMIC_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.verdict == "inconclusive"
        assert out.evidence["baseline"] == {"have": 2, "need": 3, "total": 2, "too_short": 0, "other_kind": 0}
        assert "not enough earlier work" in out.evidence["finding"]

    def test_grade_fit_counts_short_earlier_works_it_could_not_use(self, tmp_path):
        student = UserFactory()
        config = AnalyzerConfigFactory(name="text_grade_fit")
        AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence={"language": "ru", "mismatch": 0.1})
        AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence={"reason": "text too short"})
        sub = self._class_submission(level=8, user=student)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, ACADEMIC_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.evidence["baseline"] == {"have": 1, "need": 3, "total": 2, "too_short": 1, "other_kind": 0}

    def test_grade_fit_compares_sentence_length_with_earlier_work(self, tmp_path):
        student = UserFactory()
        config = AnalyzerConfigFactory(name="text_grade_fit")
        for wps, spw in ((9.0, 1.9), (10.0, 2.0), (11.0, 2.1)):
            evidence = {"language": "ru", "mismatch": 0.0, "words_per_sentence": wps, "syllables_per_word": spw}
            AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence=evidence)
        sub = self._class_submission(level=8, user=student)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, ACADEMIC_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.evidence["usual"] == {"words_per_sentence": 10.0, "syllables_per_word": 2.0, "works": 3}
        assert out.evidence["words_per_sentence"] > out.evidence["usual"]["words_per_sentence"]

    def test_grade_fit_needs_three_earlier_works_for_a_comparison(self, tmp_path):
        student = UserFactory()
        config = AnalyzerConfigFactory(name="text_grade_fit")
        evidence = {"language": "ru", "mismatch": 0.0, "words_per_sentence": 9.0, "syllables_per_word": 1.9}
        AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence=evidence)
        sub = self._class_submission(level=8, user=student)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, KID_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert "usual" not in out.evidence

    def test_grade_fit_names_bookish_words_even_when_the_text_fits(self, tmp_path):
        sub = self._class_submission(level=9)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, BOOKISH_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.verdict == "inconclusive"
        assert out.evidence["finding"].startswith("mostly fits school year 9")
        assert out.evidence["bookish_count"] == 2
        assert set(out.evidence["beyond_level_words"]) == {"байронического", "рефлексирующий"}

    def test_grade_fit_plain_text_simply_fits(self, tmp_path):
        sub = self._class_submission(level=8)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, KID_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.evidence["finding"] == "fits school year 8"
        assert "bookish_count" not in out.evidence

    def test_grade_fit_too_short_keeps_the_class_year(self, tmp_path):
        sub = self._class_submission(level=8)
        text = "Пётр Первый основал Петербург, потому что ему был нужен выход к морю. Он построил флот и открыл школы, и страна стала сильнее."
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, text), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.evidence["reason"] == "text too short"
        assert out.evidence["level"] == 8

    def test_grade_fit_flags_a_jump_above_the_students_own_level(self, tmp_path):
        student = UserFactory()
        self._grade_history(student, [0.0, 0.1, 0.05])
        sub = self._class_submission(level=8, user=student)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, ACADEMIC_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.verdict == "suspicious"
        assert out.evidence["level_label"] == "school year 8"
        assert "ai_probability" not in out.evidence

    def test_grade_fit_spares_a_student_who_always_writes_above_level(self, tmp_path):
        student = UserFactory()
        self._grade_history(student, [0.8, 0.9, 0.85])
        sub = self._class_submission(level=8, user=student)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, ACADEMIC_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.verdict == "inconclusive"
        assert "as this student's earlier work also is" in out.evidence["finding"]

    def test_grade_fit_abstains_without_level(self, tmp_path):
        sub = self._class_submission(level=None)
        out = TextGradeFitAnalyzer().analyze(_text_file(tmp_path, ACADEMIC_RU), {"submission_id": str(sub.id), "mime_type": "text/plain"})
        assert out.verdict == "inconclusive"

    def test_evasion_reports_probability_only_when_found(self, tmp_path):
        clean = TextEvasionAnalyzer().analyze(_text_file(tmp_path, KID_RU), {"mime_type": "text/plain"})
        assert clean.verdict == "inconclusive" and "ai_probability" not in clean.evidence
        dirty = TextEvasionAnalyzer().analyze(_text_file(tmp_path, KID_RU.replace("а", "a"), "d.txt"), {"mime_type": "text/plain"})
        assert dirty.verdict == "fake" and dirty.evidence["ai_probability"] >= 0.9

    def test_style_uses_the_students_earlier_class_work(self, tmp_path):
        student = UserFactory()
        config = AnalyzerConfigFactory(name="text_style")
        lang = detect_language(KID_RU)
        for variant in (KID_RU, KID_RU.replace("бабушке", "дедушке"), KID_RU.replace("щуку", "окуня")):
            earlier = self._class_submission(user=student)
            AnalysisResultFactory(submission=earlier, analyzer=config, verdict="inconclusive", evidence={"profile": {**features(variant, lang), "group": "personal"}})
        current = self._class_submission(user=student)
        out = TextStyleAnalyzer().analyze(_text_file(tmp_path, ADULT_NARRATIVE_RU), {"submission_id": str(current.id), "mime_type": "text/plain"})
        assert out.evidence["group"] == "personal"
        assert out.evidence["drift"]["history_size"] == 3
        assert out.verdict == "suspicious"

    def test_style_never_compares_different_kinds_of_text(self, tmp_path):
        student = UserFactory()
        config = AnalyzerConfigFactory(name="text_style")
        lang = detect_language(KID_RU)
        for _ in range(3):
            earlier = self._class_submission(user=student)
            AnalysisResultFactory(submission=earlier, analyzer=config, evidence={"profile": {**features(KID_RU, lang), "group": "personal"}})
        current = self._class_submission(user=student)
        out = TextStyleAnalyzer().analyze(_text_file(tmp_path, ACADEMIC_RU), {"submission_id": str(current.id), "mime_type": "text/plain"})
        assert out.evidence["group"] == "expository"
        assert out.verdict == "inconclusive" and "drift" not in out.evidence
        assert "same kind of text" in out.evidence["finding"]

    def test_style_never_compares_with_other_students(self, tmp_path):
        config = AnalyzerConfigFactory(name="text_style")
        lang = detect_language(KID_RU)
        for _ in range(3):
            AnalysisResultFactory(submission=self._class_submission(), analyzer=config, evidence={"profile": features(KID_RU, lang)})
        out = TextStyleAnalyzer().analyze(_text_file(tmp_path, ACADEMIC_RU), {"submission_id": str(self._class_submission().id), "mime_type": "text/plain"})
        assert out.verdict == "inconclusive"
        assert "drift" not in out.evidence
        assert out.evidence["baseline"] == {"have": 0, "need": 3, "total": 0, "too_short": 0, "other_kind": 0}

    def test_style_baseline_says_why_earlier_works_were_left_out(self, tmp_path):
        student = UserFactory()
        config = AnalyzerConfigFactory(name="text_style")
        lang = detect_language(KID_RU)
        AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence={"profile": {**features(KID_RU, lang), "group": "personal"}})
        AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence={"profile": {**features(ACADEMIC_RU, lang), "group": "expository"}})
        AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence={"reason": "text too short"})
        current = self._class_submission(user=student)
        out = TextStyleAnalyzer().analyze(_text_file(tmp_path, ADULT_NARRATIVE_RU), {"submission_id": str(current.id), "mime_type": "text/plain"})
        assert out.evidence["baseline"] == {"have": 1, "need": 3, "total": 3, "too_short": 1, "other_kind": 1}

    def test_style_on_a_short_answer_says_it_is_too_short(self, tmp_path):
        out = TextStyleAnalyzer().analyze(_text_file(tmp_path, "Пётр I основал Петербург в 1703 году, чтобы выйти к морю."), {"mime_type": "text/plain"})
        assert out.evidence["reason"] == "text too short"
        assert out.evidence["min_words"] == 60

    def test_style_shows_what_changed_when_it_flags(self, tmp_path):
        student = UserFactory()
        config = AnalyzerConfigFactory(name="text_style")
        lang = detect_language(KID_RU)
        for variant in (KID_RU, KID_RU.replace("бабушке", "дедушке"), KID_RU.replace("щуку", "окуня")):
            AnalysisResultFactory(submission=self._class_submission(user=student), analyzer=config, evidence={"profile": {**features(variant, lang), "group": "personal"}})
        current = self._class_submission(user=student)
        out = TextStyleAnalyzer().analyze(_text_file(tmp_path, ADULT_NARRATIVE_RU), {"submission_id": str(current.id), "mime_type": "text/plain"})
        assert out.verdict == "suspicious"
        changes = out.evidence["drift"]["changes"]
        assert changes and {"key", "now", "usual"} <= set(changes[0])

    def test_style_ignores_chatbot_phrases_that_come_from_the_assignment(self, tmp_path):
        current = self._class_submission()
        ShadowSet.objects.create(classroom=current.classroom, prompt="Как жить в современном мире без книг?", language="ru", status=ShadowSet.Status.DONE)
        text = KID_RU + " В современном мире всё быстро меняется, и чтение представляет собой отдых."
        out = TextStyleAnalyzer().analyze(_text_file(tmp_path, text), {"submission_id": str(current.id), "mime_type": "text/plain"})
        assert out.evidence["phrasebook"]["hits"] == {"представляет собой": 1}


@pytest.mark.django_db
class TestContextAggregation:
    def _result(self, sub, name, verdict, confidence, evidence=None):
        return AnalysisResultFactory(submission=sub, analyzer=AnalyzerConfigFactory(name=name, weight=2.5), verdict=verdict, confidence=confidence, evidence=evidence or {})

    def test_context_flag_turns_authentic_into_review(self):
        sub = SubmissionFactory()
        results = [
            self._result(sub, "llm_text", AnalysisResult.Verdict.AUTHENTIC, 0.8),
            self._result(sub, "text_grade_fit", AnalysisResult.Verdict.SUSPICIOUS, 0.9),
        ]
        assert aggregate(results)[1] == "needs_review"

    def test_context_alone_never_convicts(self):
        sub = SubmissionFactory()
        results = [
            self._result(sub, "text_grade_fit", AnalysisResult.Verdict.SUSPICIOUS, 1.0),
            self._result(sub, "text_style", AnalysisResult.Verdict.SUSPICIOUS, 1.0),
        ]
        assert aggregate(results)[1] == "needs_review"

    def test_weak_context_is_ignored(self):
        sub = SubmissionFactory()
        results = [
            self._result(sub, "llm_text", AnalysisResult.Verdict.AUTHENTIC, 0.8),
            self._result(sub, "text_style", AnalysisResult.Verdict.INCONCLUSIVE, 0.9),
        ]
        assert aggregate(results)[1] == "authentic"


@pytest.mark.django_db
class TestNoContentDetector:
    def test_metadata_alone_cannot_call_an_image_authentic(self):
        sub = SubmissionFactory()
        results = [
            AnalysisResultFactory(submission=sub, analyzer=AnalyzerConfigFactory(name="metadata", weight=1.5), verdict=AnalysisResult.Verdict.AUTHENTIC, confidence=0.7),
            AnalysisResultFactory(submission=sub, analyzer=AnalyzerConfigFactory(name="ela", weight=0.75), verdict=AnalysisResult.Verdict.INCONCLUSIVE, confidence=0.5),
            AnalysisResultFactory(submission=sub, analyzer=AnalyzerConfigFactory(name="community_forensics", weight=3.5), verdict=AnalysisResult.Verdict.ERROR, confidence=0.0),
        ]
        assert aggregate(results) == (0.5, "inconclusive")

    def test_metadata_ai_declaration_still_counts(self):
        sub = SubmissionFactory()
        results = [
            AnalysisResultFactory(submission=sub, analyzer=AnalyzerConfigFactory(name="metadata", weight=1.5), verdict=AnalysisResult.Verdict.FAKE, confidence=0.95, evidence={"ai_probability": 0.95}),
            AnalysisResultFactory(submission=sub, analyzer=AnalyzerConfigFactory(name="community_forensics", weight=3.5), verdict=AnalysisResult.Verdict.ERROR, confidence=0.0),
        ]
        assert aggregate(results)[1] != "authentic"
